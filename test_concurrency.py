from concurrent.futures import Future, ThreadPoolExecutor, as_completed
from threading import Barrier

from openpyxl import load_workbook

from ai_agent import data_access
from ai_agent.tools import execute_tool

CLIENT_IDS = (
    "+18095550001",
    "+18095550002",
    "+18095550003",
    "+18095550004",
    "+18095550005",
)
TOOL_NAME = "save_order"


def read_orders() -> tuple[tuple[object, ...], list[tuple[object, ...]]]:
    if not data_access.ORDERS_FILE.exists():
        return (), []

    workbook = load_workbook(
        data_access.ORDERS_FILE,
        read_only=True,
        data_only=True,
    )
    try:
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(values_only=True))
        headers = tuple(rows[0]) if rows else ()
        orders = [
            tuple(row)
            for row in rows[1:]
            if any(value is not None for value in row)
        ]
        return headers, orders
    finally:
        workbook.close()


def save_test_order(client_id: str, start_gate: Barrier) -> str:
    arguments = {
        "customer_name": f"Concurrency Test {client_id[-4:]}",
        "items": [{"product": "cookie", "quantity": 1}],
    }
    start_gate.wait()
    return execute_tool(TOOL_NAME, arguments, client_id)


def main() -> None:
    _, existing_orders = read_orders()
    existing_order_count = len(existing_orders)
    start_gate = Barrier(len(CLIENT_IDS))
    successful_clients: list[str] = []
    completed_count = 0

    with ThreadPoolExecutor(max_workers=len(CLIENT_IDS)) as executor:
        futures: dict[Future[str], str] = {
            executor.submit(save_test_order, client_id, start_gate): client_id
            for client_id in CLIENT_IDS
        }

        for future in as_completed(futures):
            client_id = futures[future]
            try:
                result = future.result()
            except Exception as error:
                completed_count += 1
                print(f"{client_id}: EXCEPTION: {error}")
                continue

            completed_count += 1
            if "registered successfully" in result:
                successful_clients.append(client_id)
                print(f"{client_id}: SUCCESS: {result}")
            else:
                print(f"{client_id}: FAILED: {result}")

    print(f"Operations completed: {completed_count}/{len(CLIENT_IDS)}")

    try:
        headers, orders = read_orders()
    except Exception as error:
        print(f"Workbook validation: FAILED to read orders.xlsx: {error}")
        return

    if headers != data_access.ORDER_HEADERS:
        print(f"Workbook validation: FAILED; unexpected headers: {headers}")
        return

    contact_column = headers.index("contact")
    saved_contacts = [row[contact_column] for row in orders]
    missing_clients = [
        client_id
        for client_id in successful_clients
        if saved_contacts.count(client_id) != 1
    ]
    expected_order_count = existing_order_count + len(successful_clients)
    if len(orders) != expected_order_count:
        print(
            "Workbook validation: FAILED; "
            f"expected {expected_order_count} orders, found {len(orders)}."
        )
    elif missing_clients:
        print(
            "Workbook validation: FAILED; successful results were not saved "
            f"exactly once for: {', '.join(missing_clients)}."
        )
    else:
        print(
            "Workbook validation: PASSED; each successful result appears "
            "exactly once with aligned headers."
        )


if __name__ == "__main__":
    main()
