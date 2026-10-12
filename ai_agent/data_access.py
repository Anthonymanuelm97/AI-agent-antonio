from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = PROJECT_ROOT / "data"
BUSINESS_DATA_FILE = DATA_DIR / "cafe_santo_domingo_business_data.xlsx"

_REQUIRED_COLUMNS = {
    "Products": ("product", "price (RD$)", "category"),
    "Information": ("topic", "detail"),
}


def load_business_data() -> dict[str, Any]:
    if not BUSINESS_DATA_FILE.is_file():
        raise FileNotFoundError(
            f"Business data workbook not found: {BUSINESS_DATA_FILE}"
        )

    try:
        workbook = load_workbook(
            BUSINESS_DATA_FILE,
            read_only=True,
            data_only=True,
        )
    except (OSError, BadZipFile, InvalidFileException) as error:
        raise ValueError(
            f"Unable to open business data workbook "
            f"{BUSINESS_DATA_FILE}: {error}"
        ) from error

    try:
        missing_sheets = set(_REQUIRED_COLUMNS) - set(workbook.sheetnames)
        if missing_sheets:
            raise ValueError(
                "Business data workbook is missing required sheet(s): "
                + ", ".join(sorted(missing_sheets))
            )

        sheets: dict[str, list[dict[str, Any]]] = {}
        for sheet_name, required_columns in _REQUIRED_COLUMNS.items():
            worksheet = workbook[sheet_name]
            headers = next(
                worksheet.iter_rows(
                    min_row=1,
                    max_row=1,
                    values_only=True,
                ),
                (),
            )
            missing_columns = set(required_columns) - set(headers)
            if missing_columns:
                raise ValueError(
                    f"Sheet '{sheet_name}' is missing required column(s): "
                    + ", ".join(sorted(missing_columns))
                )

            sheets[sheet_name] = [
                dict(zip(headers, values))
                for values in worksheet.iter_rows(
                    min_row=2,
                    values_only=True,
                )
                if any(value is not None for value in values)
            ]
    finally:
        workbook.close()

    price_column = _REQUIRED_COLUMNS["Products"][1]
    currency = price_column.removeprefix("price (").removesuffix(")")
    product_rows = [
        product
        for product in sheets["Products"]
        if product.get("product") and product.get("category")
    ]
    products = [
        {
            "name": product["product"],
            "price": product["price (RD$)"],
            "category": product["category"],
        }
        for product in product_rows
    ]

    return {
        "products": products,
        "information": sheets["Information"],
        "currency": currency,
    }
